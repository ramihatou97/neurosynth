"""DOCX document parser using python-docx."""

import asyncio
from pathlib import Path
from typing import Any

from rich.console import Console

from neurosynth.models.document import Document, DocumentFormat
from neurosynth.parsers.base import BaseParser

console = Console()


class DOCXParser(BaseParser):
    """Parser for DOCX documents using python-docx."""

    supported_formats = [DocumentFormat.DOCX]

    async def parse(self, path: Path) -> Document:
        """Parse a DOCX document."""
        if not path.exists():
            raise FileNotFoundError(f"DOCX not found: {path}")

        # Extract metadata
        metadata = await self.extract_metadata(path)
        source = self._create_source(path, metadata)

        # Extract text
        raw_text = await self.extract_text(path)

        # Create document
        doc = Document(
            source=source,
            raw_text=raw_text,
            is_parsed=True,
        )

        # Extract structure
        doc.toc = await self._extract_headings(path)

        return doc

    async def extract_text(self, path: Path) -> str:
        """Extract text from DOCX."""
        loop = asyncio.get_event_loop()

        def extract():
            from docx import Document as DocxDocument

            doc = DocxDocument(str(path))
            text_parts = []

            for para in doc.paragraphs:
                text = para.text.strip()
                if text:
                    # Check if it's a heading
                    if para.style and para.style.name.startswith("Heading"):
                        level = self._get_heading_level(para.style.name)
                        prefix = "#" * level
                        text_parts.append(f"\n{prefix} {text}\n")
                    else:
                        text_parts.append(text)

            # Also extract from tables
            for table in doc.tables:
                table_text = self._extract_table_text(table)
                if table_text:
                    text_parts.append(f"\n[Table]\n{table_text}\n")

            return "\n\n".join(text_parts)

        return await loop.run_in_executor(None, extract)

    async def extract_metadata(self, path: Path) -> dict[str, Any]:
        """Extract metadata from DOCX."""
        loop = asyncio.get_event_loop()

        def extract():
            from docx import Document as DocxDocument

            doc = DocxDocument(str(path))
            props = doc.core_properties

            # Get authors
            authors = []
            if props.author:
                authors = [a.strip() for a in props.author.split(",")]

            # Get year from dates
            year = None
            if props.created:
                year = props.created.year
            elif props.modified:
                year = props.modified.year

            return {
                "title": props.title or "",
                "authors": authors,
                "year": year,
                "publisher": props.category or "",
            }

        return await loop.run_in_executor(None, extract)

    async def _extract_headings(self, path: Path) -> list[dict[str, Any]]:
        """Extract document headings as TOC."""
        loop = asyncio.get_event_loop()

        def extract():
            from docx import Document as DocxDocument

            doc = DocxDocument(str(path))
            toc = []

            for para in doc.paragraphs:
                if para.style and para.style.name.startswith("Heading"):
                    level = self._get_heading_level(para.style.name)
                    title = para.text.strip()
                    if title:
                        toc.append(
                            {
                                "level": level,
                                "title": title,
                            }
                        )

            return toc

        return await loop.run_in_executor(None, extract)

    def _get_heading_level(self, style_name: str) -> int:
        """Extract heading level from style name."""
        import re

        match = re.search(r"Heading\s*(\d+)", style_name)
        if match:
            return int(match.group(1))
        return 1

    def _extract_table_text(self, table) -> str:
        """Extract text from a table."""
        rows = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            rows.append(" | ".join(cells))
        return "\n".join(rows)

    async def extract_sections(self, path: Path) -> list[dict[str, Any]]:
        """Extract document sections based on headings."""
        loop = asyncio.get_event_loop()

        def extract():
            from docx import Document as DocxDocument

            doc = DocxDocument(str(path))
            sections = []
            current_section = None
            current_content = []

            for para in doc.paragraphs:
                text = para.text.strip()

                # Check if it's a heading
                if para.style and para.style.name.startswith("Heading"):
                    # Save previous section
                    if current_section:
                        sections.append(
                            {
                                "title": current_section,
                                "content": "\n\n".join(current_content),
                            }
                        )

                    # Start new section
                    current_section = text
                    current_content = []
                elif text and current_section:
                    current_content.append(text)

            # Save last section
            if current_section:
                sections.append(
                    {
                        "title": current_section,
                        "content": "\n\n".join(current_content),
                    }
                )

            return sections

        return await loop.run_in_executor(None, extract)
